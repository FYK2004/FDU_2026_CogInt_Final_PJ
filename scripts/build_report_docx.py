"""Render docs/项目报告.md into a submission-ready Word document.

目标：正文 <= 8 页（图表、参考文献、附录另算），因此版式压紧：
A4、页边距 2.3/2.5cm、正文 10.5pt、行距 1.15、段后 4pt；图统一 13cm 宽。

用法（需要 python-docx）：
    python scripts/build_report_docx.py
    python scripts/build_report_docx.py docs/项目原理说明.md docs/项目原理说明.docx

不传参数时渲染期末报告；传入源文件 / 输出文件即可用同一套版式渲染其他 Markdown
（例如分步原理说明），保证「单一来源、两种形态」且版式一致。
"""

from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "docs", "项目报告.md")
OUT = os.path.join(ROOT, "docs", "第XX小组-项目期末报告.docx")

LATIN = "Times New Roman"
CJK_HEAD = "微软雅黑"
CJK_BODY = "宋体"
IMG_WIDTH_CM = 13.0


def set_run_font(run, size=None, bold=None, latin=LATIN, cjk=CJK_BODY,
                 color=None, mono=False):
    if mono:
        latin = cjk = "Consolas"
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cjk)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def setup(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.3)
    sec.bottom_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_BODY)
    pf = normal.paragraph_format
    pf.line_spacing = 1.15
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)


INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")


def unescape(text):
    return text.replace("\\_", "_").replace("\\*", "*").replace("\\`", "`")


def _add_plain(par, chunk, size, bold, color):
    for i, piece in enumerate(chunk.split("`")):
        if not piece:
            continue
        run = par.add_run(piece)
        set_run_font(run, size=size, bold=bold, color=color, mono=(i % 2 == 1))


def add_inline(par, text, size=10.5, base_bold=False, color=None):
    text = unescape(text)
    pos = 0
    for m in INLINE_BOLD.finditer(text):
        if m.start() > pos:
            _add_plain(par, text[pos:m.start()], size, base_bold, color)
        _add_plain(par, m.group(1), size, True, color)
        pos = m.end()
    if pos < len(text):
        _add_plain(par, text[pos:], size, base_bold, color)


def add_heading(doc, text, level):
    sizes = {1: 17.0, 2: 13.5, 3: 11.5}
    size = sizes[level]
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(0 if level == 1 else 9)
    par.paragraph_format.space_after = Pt(6)
    if level == 1:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(par, text, size=size, base_bold=True)
    for run in par.runs:
        set_run_font(run, size=size, bold=True, cjk=CJK_HEAD)
    return par


def add_para(doc, text, indent_first=True):
    par = doc.add_paragraph()
    par.paragraph_format.first_line_indent = Cm(0.74) if indent_first else Cm(0)
    add_inline(par, text)
    return par


def add_bullet(doc, text, level=0):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.74 + 0.6 * level)
    par.paragraph_format.first_line_indent = Cm(-0.37)
    par.paragraph_format.space_after = Pt(2)
    run = par.add_run("• " if level == 0 else "– ")
    set_run_font(run, size=10.5)
    add_inline(par, text)
    return par


def add_mono(doc, text):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(1.0)
    par.paragraph_format.space_after = Pt(2)
    run = par.add_run(unescape(text))
    set_run_font(run, size=9.5, mono=True)
    return par


def add_note(doc, text):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.74)
    add_inline(par, text, size=9.5, color=RGBColor(0x55, 0x55, 0x55))
    return par


def add_image(doc, caption, path):
    full = path if os.path.isabs(path) else os.path.join(ROOT, path)
    if not os.path.exists(full):
        add_note(doc, "[缺图] " + path)
        return
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(2)
    par.add_run().add_picture(full, width=Cm(IMG_WIDTH_CM))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    add_inline(cap, caption, size=9, color=RGBColor(0x33, 0x33, 0x33))


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, cell_text in enumerate(header):
        cell = table.cell(0, j)
        cell.text = ""
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(1)
        add_inline(par, cell_text, size=9, base_bold=True)
        for run in par.runs:
            set_run_font(run, size=9, bold=True)
    for i, row in enumerate(body, start=1):
        for j in range(len(header)):
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            add_inline(par, row[j] if j < len(row) else "", size=9)
            for run in par.runs:
                set_run_font(run, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


IMG = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")


def split_table_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def is_sep(line):
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c) for c in cells)


def main():
    global SRC, OUT
    if len(sys.argv) >= 2:
        SRC = sys.argv[1] if os.path.isabs(sys.argv[1]) else os.path.join(ROOT, sys.argv[1])
    if len(sys.argv) >= 3:
        OUT = sys.argv[2] if os.path.isabs(sys.argv[2]) else os.path.join(ROOT, sys.argv[2])

    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    setup(doc)

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if re.fullmatch(r"-{3,}", stripped):
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            add_heading(doc, m.group(2), min(len(m.group(1)), 3))
            i += 1
            continue

        m = IMG.match(stripped)
        if m:
            add_image(doc, m.group(1), m.group(2))
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_sep(lines[i + 1]):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_sep(lines[i]):
                    block.append(split_table_row(lines[i]))
                i += 1
            add_table(doc, block)
            continue

        if stripped.startswith(">"):
            add_note(doc, stripped.lstrip("> ").strip())
            i += 1
            continue

        if line.startswith("    ") and not stripped.startswith("- "):
            block = []
            while i < len(lines) and lines[i].startswith("    "):
                block.append(lines[i][4:])
                i += 1
            for b in block:
                add_mono(doc, b)
            continue

        if stripped.startswith("- "):
            add_bullet(doc, stripped[2:].strip(), (len(line) - len(line.lstrip())) // 2)
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.74)
            par.paragraph_format.first_line_indent = Cm(-0.74)
            add_inline(par, m.group(1) + ". " + m.group(2))
            i += 1
            continue

        add_para(doc, stripped)
        i += 1

    doc.save(OUT)
    print("wrote " + OUT)


if __name__ == "__main__":
    sys.exit(main())
